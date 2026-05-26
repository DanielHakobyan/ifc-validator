import io
import os
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = os.path.join(BASE_DIR, '..', 'fonts')

pdfmetrics.registerFont(TTFont('DejaVu', os.path.join(FONT_DIR, 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DejaVu-Bold', os.path.join(FONT_DIR, 'DejaVuSans.ttf')))

def generate_pdf(summary: dict, issues: list) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm,
                             topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        fontName='DejaVu-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1565C0'),
        spaceAfter=6,
    )
    normal_style = ParagraphStyle(
        'CustomNormal',
        fontName='DejaVu',
        fontSize=10,
        spaceAfter=4,
    )
    heading2_style = ParagraphStyle(
        'CustomHeading2',
        fontName='DejaVu-Bold',
        fontSize=13,
        textColor=colors.HexColor('#1565C0'),
        spaceBefore=10,
        spaceAfter=6,
    )
    cell_style = ParagraphStyle(
        'CellStyle',
        fontName='DejaVu',
        fontSize=8,
        leading=10,
    )

    story.append(Paragraph("ТИМ Инспектор — Отчёт о проверке", title_style))
    story.append(Paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}", normal_style))
    story.append(Spacer(1, 0.5*cm))

    # Summary table
    summary_data = [
        ['Показатель', 'Значение'],
        ['Всего элементов', str(summary.get('total', 0))],
        ['Ошибок', str(summary.get('errors', 0))],
        ['Предупреждений', str(summary.get('warnings', 0))],
        ['Прошло проверку', str(summary.get('passed', 0))],
    ]
    summary_table = Table(summary_data, colWidths=[10*cm, 6*cm])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565C0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), 'DejaVu'),
        ('FONTNAME', (0, 0), (-1, 0), 'DejaVu-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
    ]))

    story.append(Paragraph("Сводка результатов", heading2_style))
    story.append(summary_table)
    story.append(Spacer(1, 0.5*cm))

    # Issues table
    story.append(Paragraph("Список проблем", heading2_style))
    if issues:
        issues_data = [['#', 'GUID', 'Тип', 'Статус', 'Описание', 'Местоположение']]
        for issue in issues:
            issues_data.append([
                str(issue.get('id', '')),
                str(issue.get('guid', ''))[:10],
                str(issue.get('type', '')),
                str(issue.get('severity', '')),
                Paragraph(str(issue.get('description', '')), cell_style),
                Paragraph(str(issue.get('location', '')), cell_style),
            ])
        issues_table = Table(issues_data, colWidths=[0.8*cm, 2.2*cm, 2.5*cm, 2.5*cm, 5*cm, 4*cm])
        issues_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565C0')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, -1), 'DejaVu'),
            ('FONTNAME', (0, 0), (-1, 0), 'DejaVu-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF9F9')]),
        ]))
        story.append(issues_table)

    doc.build(story)
    return buffer.getvalue()


def generate_excel(summary: dict, issues: list) -> bytes:
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Сводка"

    blue_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    ws_summary['A1'] = 'ТИМ Инспектор — Отчёт о проверке'
    ws_summary['A1'].font = Font(bold=True, size=14, color="1565C0")
    ws_summary['A2'] = f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws_summary.append([])
    ws_summary.append(['Показатель', 'Значение'])
    ws_summary['A4'].fill = blue_fill; ws_summary['A4'].font = header_font
    ws_summary['B4'].fill = blue_fill; ws_summary['B4'].font = header_font
    for key, val in [('Всего элементов', summary.get('total', 0)),
                     ('Ошибок', summary.get('errors', 0)),
                     ('Предупреждений', summary.get('warnings', 0)),
                     ('Прошло проверку', summary.get('passed', 0))]:
        ws_summary.append([key, val])

    ws_summary.column_dimensions['A'].width = 25
    ws_summary.column_dimensions['B'].width = 15

    ws_issues = wb.create_sheet("Проблемы")
    headers = ['#', 'GUID', 'Тип', 'Статус', 'Описание', 'Местоположение']
    ws_issues.append(headers)
    for i, cell in enumerate(ws_issues[1], 1):
        cell.fill = blue_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    severity_fills = {
        'ERROR': PatternFill(start_color="FFF5F5", end_color="FFF5F5", fill_type="solid"),
        'WARNING': PatternFill(start_color="FFFBF0", end_color="FFFBF0", fill_type="solid"),
        'OK': PatternFill(start_color="F5FFF5", end_color="F5FFF5", fill_type="solid"),
    }

    for issue in issues:
        row = [
            issue.get('id', ''), issue.get('guid', ''), issue.get('type', ''),
            issue.get('severity', ''), issue.get('description', ''), issue.get('location', '')
        ]
        ws_issues.append(row)
        sev = issue.get('severity', '')
        fill = severity_fills.get(sev)
        if fill:
            for cell in ws_issues[ws_issues.max_row]:
                cell.fill = fill
                cell.border = border

    col_widths = [5, 15, 15, 18, 55, 30]
    for i, w in enumerate(col_widths, 1):
        ws_issues.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_bcf(issues: list) -> str:
    today = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    topics = ""
    for issue in issues:
        if issue.get("severity") in ("ERROR", "WARNING"):
            topics += f"""  <Topic Guid="{issue.get('guid', '')}-BCF" TopicType="Error" TopicStatus="Open">
    <Title>{issue.get('description', '')}</Title>
    <Description>Тип: {issue.get('type', '')} | Статус: {issue.get('severity', '')} | Местоположение: {issue.get('location', '')}</Description>
    <CreationDate>{today}</CreationDate>
    <CreationAuthor>TIM Inspector</CreationAuthor>
  </Topic>\n"""

    return f"""<?xml version="1.0" encoding="UTF-8"?>
<BcfRootProject>
  <ProjectExtension>
    <Project ProjectId="TIM-{datetime.now().strftime('%Y%m%d')}">
      <Name>TIM Inspector Report</Name>
    </Project>
  </ProjectExtension>
  <Markup>
    <Header>
      <File IfcProject="TIM Inspector">
        <Filename>model.ifc</Filename>
        <Date>{today}</Date>
      </File>
    </Header>
{topics}  </Markup>
</BcfRootProject>"""